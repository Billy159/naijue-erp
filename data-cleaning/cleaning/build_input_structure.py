"""Freeze workbook paths, sheet order, raw headers and observed value types."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def build(input_dir: Path, audit_profiles: Path) -> dict:
    profiles = json.loads(audit_profiles.read_text())
    profile_index = {
        (Path(item["file"]).relative_to("T+数据").as_posix(), item["sheet"]): item
        for item in profiles
    }
    files = []
    sheets = []
    for path in sorted(input_dir.rglob("*")):
        if not path.is_file() or path.name == ".DS_Store" or path.name.startswith(".~"):
            continue
        rel = path.relative_to(input_dir).as_posix()
        if path.suffix not in (".xlsx", ".bof"):
            raise ValueError(f"Unexpected source file: {rel}")
        files.append({"path": rel, "kind": "workbook" if path.suffix == ".xlsx" else "print_template"})
        if path.suffix != ".xlsx":
            continue
        book = load_workbook(path, read_only=True, data_only=False)
        for sheet in book:
            info = profile_index[(rel, sheet.title)]
            header_row = info["header_row"]
            if rel == "库存状况表.xlsx":
                header_row = 7
            headers = list(next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))) if header_row else []
            sheets.append({
                "file": rel, "sheet": sheet.title, "sheetIndex": book.sheetnames.index(sheet.title),
                "headerRow": header_row, "dataStartRow": 9 if rel == "库存状况表.xlsx" else header_row + 1,
                "headers": headers,
                "maxColumn": sheet.max_column,
                "observedTypes": [sorted(field["types"]) for field in info["fields"]],
            })
        book.close()
    payload = {"inputStructureVersion": "TPLUS-STRUCTURE-v1", "files": files, "sheets": sheets}
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    payload["fingerprintSha256"] = hashlib.sha256(canonical.encode()).hexdigest()
    return payload


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--audit-profiles", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    result = build(args.input, args.audit_profiles)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str) + "\n")
    print(json.dumps({"files": len(result["files"]), "sheets": len(result["sheets"]), "fingerprint": result["fingerprintSha256"]}))
