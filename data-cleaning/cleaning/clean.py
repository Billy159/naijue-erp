"""Strict, repeatable T+ Excel cleaner. Original files are opened read-only."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
import hashlib
from itertools import zip_longest
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
from xml.sax.saxutils import escape
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string


BASE = Path(__file__).resolve().parent
FIELDNAMES = ["issueId", "domain", "sourceFile", "sheet", "row", "entity", "problem", "candidateA", "candidateB", "recommendation", "confidence", "userDecision", "finalAction"]


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def json_write(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n")


def issue_id(*parts: object) -> str:
    return "TPCI-" + hashlib.sha256(json.dumps(parts, ensure_ascii=False, default=str).encode()).hexdigest()[:12]


def review(rel: str, sheet: str, row: int, entity: str, problem: str, a: str, b: str, recommendation: str, domain: str = "基础资料", confidence: str = "MED") -> dict:
    return dict(zip(FIELDNAMES, [issue_id(rel, sheet, row, problem), domain, "T+数据/" + rel, sheet, row, entity, problem, a, b, recommendation, confidence, "", ""]))


def validate_contract(source: Path, structure: dict) -> tuple[list[str], list[dict]]:
    expected = {item["path"] for item in structure["files"]}
    actual = {p.relative_to(source).as_posix() for p in source.rglob("*") if p.is_file() and p.name != ".DS_Store" and not p.name.startswith(".~")}
    drift = []
    for path in sorted(expected - actual):
        drift.append({"kind": "missing_file", "path": path})
    for path in sorted(actual - expected):
        drift.append({"kind": "extra_file", "path": path})
    by_file = defaultdict(list)
    for item in structure["sheets"]:
        by_file[item["file"]].append(item)
    for rel, sheets in by_file.items():
        if rel not in actual:
            continue
        book = load_workbook(source / rel, read_only=True, data_only=False)
        expected_names = [item["sheet"] for item in sheets]
        if book.sheetnames != expected_names:
            drift.append({"kind": "sheet_order_or_name", "path": rel, "expected": expected_names, "actual": book.sheetnames})
            book.close()
            continue
        for sheet, info in zip(book, sheets):
            if sheet.max_column != info["maxColumn"]:
                drift.append({"kind": "column_count", "path": rel, "sheet": sheet.title, "expected": info["maxColumn"], "actual": sheet.max_column})
            if info["headerRow"]:
                headers = list(next(sheet.iter_rows(min_row=info["headerRow"], max_row=info["headerRow"], values_only=True)))
                if headers != info["headers"]:
                    drift.append({"kind": "headers", "path": rel, "sheet": sheet.title, "expected": info["headers"], "actual": headers})
            accepted = [set(types) for types in info["observedTypes"]]
            found_new = set()
            for row in sheet.iter_rows(min_row=info["dataStartRow"], values_only=True):
                for column, value in enumerate(row):
                    if value is None or value == "":
                        continue
                    if column >= len(accepted) or type(value).__name__ not in accepted[column]:
                        found_new.add((column + 1, type(value).__name__))
            for column, value_type in sorted(found_new):
                drift.append({"kind": "column_type", "path": rel, "sheet": sheet.title, "column": get_column_letter(column), "newType": value_type})
        book.close()
    return sorted(actual), drift


def collect(source: Path, structure: dict, manifest: dict, mapping: Path, batch: str) -> tuple[dict, list[dict], list[dict], dict]:
    edits = defaultdict(lambda: defaultdict(dict))
    changes = []
    queue = []
    rules = {item["ruleId"]: item for item in manifest["rules"]}
    if len(rules) != len(manifest["rules"]):
        raise ValueError("duplicate cleaning rule IDs")
    supported = {"TPC-001": ("AUTO_SAFE", "rstrip_chars"), "TPC-002": ("MANUAL_REVIEW", "none"), "TPC-007": ("MANUAL_REVIEW", "none"), "TPC-013": ("AUTO_WITH_ASSERTION", "assert_existing_category_projection")}
    for rule_id, (severity, operation) in supported.items():
        if rules[rule_id]["severity"] != severity or rules[rule_id]["transform"]["operation"] != operation:
            raise ValueError(f"{rule_id} manifest rule handler does not match severity or operation")
    for item in manifest["rules"]:
        if item["ruleId"] not in supported and item["transform"]["operation"] not in ("none", "preserve_existing_quarantine"):
            raise ValueError(f"unsupported transform in {item['ruleId']}")
    rule = rules["TPC-001"]
    target = next(x for x in structure["sheets"] if x["file"] == rule["applicableFilePatterns"][0] and x["sheet"] == rule["sheet"])
    description_col = target["headers"].index(rule["field"]) + 1
    category_rule = rules["TPC-013"]
    if category_rule["applicableFilePatterns"] != rule["applicableFilePatterns"] or category_rule["sheet"] != rule["sheet"]:
        raise ValueError("TPC-013 category assertion must apply to the cleaned product sheet")
    product_code_col = target["headers"].index(category_rule["field"]) + 1
    product_codes = set()
    book = load_workbook(source / target["file"], read_only=True, data_only=False)
    sheet = book[rule["sheet"]]
    for row_no, values in enumerate(sheet.iter_rows(values_only=True), 1):
        if row_no < target["dataStartRow"]:
            continue
        if values and values[product_code_col - 1] not in (None, ""):
            product_codes.add(str(values[product_code_col - 1]))
        value = values[description_col - 1] if len(values) >= description_col else None
        if not isinstance(value, str):
            continue
        after = value.rstrip(rule["transform"]["characters"])
        if after != value:
            if not after or value[:len(after)] != after:
                raise AssertionError("TPC-001 assertion failed")
            coord = f"{get_column_letter(description_col)}{row_no}"
            edits[target["file"]][rule["sheet"]][coord] = (value, after)
            changes.append({
                "rawFile": "T+数据/" + target["file"], "sheet": rule["sheet"],
                "row": row_no, "field": rule["field"], "cell": coord,
                "oldValue": value, "newValue": after, "ruleId": rule["ruleId"],
                "cleaningBatch": batch, "timestamp": datetime.now(timezone.utc).isoformat(),
            })
    book.close()
    assignments = json.loads(mapping.read_text())
    mapping_codes = [str(x["code"]) for x in assignments]
    duplicates = len(mapping_codes) - len(set(mapping_codes))
    unmatched = sorted(product_codes - set(mapping_codes))
    for code in unmatched:
        queue.append(review(target["file"], rule["sheet"], 0, code, "TPC-013: 商品缺冻结分类裁决", code, "待用户裁决分类", "保持原值，暂停该商品转换", confidence="HIGH"))
    category_assertion = {"productCount": len(product_codes), "mappingCount": len(mapping_codes), "duplicateMappings": duplicates, "unmatchedProducts": len(unmatched), "passed": duplicates == 0 and not unmatched}

    manual_rule = rules["TPC-002"]
    validation_info = next(x for x in structure["sheets"] if x["file"] == manual_rule["applicableFilePatterns"][0] and x["sheet"] == manual_rule["sheet"])
    book = load_workbook(source / validation_info["file"], read_only=True, data_only=False)
    for row_no, values in enumerate(book[manual_rule["sheet"]].iter_rows(values_only=True), 1):
        column = column_index_from_string(manual_rule["field"].rsplit("_", 1)[-1]) - 1
        value = values[column] if len(values) > column else None
        if isinstance(value, str) and value != value.strip():
            queue.append(review(validation_info["file"], manual_rule["sheet"], row_no, value.strip(), "TPC-002: 校验列表尾空格", value, value.strip(), "确认该列表项是否可统一后再修订规则"))
    book.close()

    employee_rule = rules["TPC-007"]
    employee_info = next(x for x in structure["sheets"] if x["file"] == employee_rule["applicableFilePatterns"][0] and x["sheet"] == employee_rule["sheet"])
    name_col = employee_info["headers"].index(employee_rule["field"])
    code_col = employee_info["headers"].index("员工编码")
    book = load_workbook(source / employee_info["file"], read_only=True, data_only=False)
    names = defaultdict(list)
    for row_no, values in enumerate(book[employee_rule["sheet"]].iter_rows(values_only=True), 1):
        if row_no < employee_info["dataStartRow"] or len(values) <= max(name_col, code_col) or not values[name_col]:
            continue
        names[str(values[name_col])].append((row_no, str(values[code_col])))
    for name, entries in names.items():
        if len(entries) > 1 and len({code for _, code in entries}) > 1:
            for row_no, code in entries:
                queue.append(review(employee_info["file"], employee_rule["sheet"], row_no, code, "TPC-007: 同名不同员工编码", f"{name} / {code}", ",".join(x[1] for x in entries), "保留全部员工身份，人工判断是否仅显示同名"))
    book.close()
    return edits, changes, queue, category_assertion


def patch_workbook(raw: Path, output: Path, patch_by_sheet: dict, structure: dict, rel: str) -> None:
    sheet_info = {x["sheet"]: x for x in structure["sheets"] if x["file"] == rel}
    patched_members = {}
    for name, cells in patch_by_sheet.items():
        member = f"xl/worksheets/sheet{sheet_info[name]['sheetIndex'] + 1}.xml"
        patched_members[member] = cells
    with zipfile.ZipFile(raw, "r") as inp, zipfile.ZipFile(output, "w") as out:
        if not set(patched_members) <= set(inp.namelist()):
            raise AssertionError("worksheet ZIP member missing")
        for info in inp.infolist():
            content = inp.read(info.filename)
            if info.filename in patched_members:
                for coord, (old, new) in patched_members[info.filename].items():
                    pattern = rb"<c\b(?=[^>]*\br=\"" + re.escape(coord.encode()) + rb"\")[^>]*>.*?</c>"
                    matches = list(re.finditer(pattern, content, flags=re.DOTALL))
                    if len(matches) != 1:
                        raise AssertionError(f"expected one cell {rel}:{coord}, found {len(matches)}")
                    original = matches[0].group()
                    opening = original[:original.index(b">") + 1]
                    if b"<f" in original or b't="s"' not in opening:
                        raise AssertionError(f"unsupported source cell {rel}:{coord}")
                    replacement_open = opening.replace(b't="s"', b't="inlineStr"')
                    replacement = replacement_open + b'<is><t xml:space="preserve">' + escape(new).encode() + b"</t></is></c>"
                    content = content[:matches[0].start()] + replacement + content[matches[0].end():]
            out.writestr(info, content)


def verify_output(source: Path, cleaned: Path, edits: dict, raw_hashes: dict) -> dict:
    cleaned_hashes = {}
    for rel, before_sha in raw_hashes.items():
        output = cleaned / rel
        if not output.exists():
            raise AssertionError(f"missing cleaned file: {rel}")
        cleaned_hashes[rel] = digest(output)
        if rel not in edits and cleaned_hashes[rel] != before_sha:
            raise AssertionError(f"silent byte mutation: {rel}")
        if rel not in edits:
            continue
        with zipfile.ZipFile(source / rel) as raw_zip, zipfile.ZipFile(output) as clean_zip:
            if raw_zip.namelist() != clean_zip.namelist():
                raise AssertionError(f"ZIP member drift: {rel}")
            changed_members = [name for name in raw_zip.namelist() if raw_zip.read(name) != clean_zip.read(name)]
            sheet_book = load_workbook(source / rel, read_only=True)
            expected_members = [f"xl/worksheets/sheet{index + 1}.xml" for index, name in enumerate(sheet_book.sheetnames) if name in edits[rel]]
            sheet_book.close()
            if changed_members != expected_members:
                raise AssertionError(f"unexpected XLSX member changes: {rel}: {changed_members}")
        raw_book = load_workbook(source / rel, read_only=True, data_only=False)
        clean_book = load_workbook(output, read_only=True, data_only=False)
        actual_diffs = {}
        for raw_sheet, clean_sheet in zip(raw_book, clean_book):
            if raw_sheet.title != clean_sheet.title:
                raise AssertionError(f"sheet name drift: {rel}")
            if (raw_sheet.max_row, raw_sheet.max_column) != (clean_sheet.max_row, clean_sheet.max_column):
                raise AssertionError(f"sheet dimension drift: {rel}:{raw_sheet.title}")
            for row_num, (old_row, new_row) in enumerate(zip_longest(raw_sheet.iter_rows(), clean_sheet.iter_rows()), 1):
                if old_row is None or new_row is None or len(old_row) != len(new_row):
                    raise AssertionError(f"sheet row drift: {rel}:{raw_sheet.title}:{row_num}")
                for column, (old_cell, new_cell) in enumerate(zip(old_row, new_row), 1):
                    if old_cell.value != new_cell.value:
                        actual_diffs[(raw_sheet.title, f"{get_column_letter(column)}{row_num}")] = (old_cell.value, new_cell.value)
                    if getattr(old_cell, "_style_id", None) != getattr(new_cell, "_style_id", None) or old_cell.number_format != new_cell.number_format:
                        raise AssertionError(f"style drift: {rel}:{raw_sheet.title}:{old_cell.coordinate}")
        expected_diffs = {(sheet, coord): change for sheet, cells in edits[rel].items() for coord, change in cells.items()}
        if actual_diffs != expected_diffs:
            raise AssertionError(f"silent cell mutation: {rel}: {len(actual_diffs)} vs {len(expected_diffs)}")
        raw_book.close()
        clean_book.close()
    if any(digest(source / rel) != sha for rel, sha in raw_hashes.items()):
        raise AssertionError("raw file changed during cleaning")
    return cleaned_hashes


def render_profiles(run: Path, changes: list[dict], queue: list[dict], category: dict) -> None:
    before = json.loads((run / "before-scan.json").read_text())
    after = json.loads((run / "after-scan.json").read_text())
    def table(profile: dict) -> str:
        return "\n".join(f"| {kind} | {count:,} |" for kind, count in profile["counts"].items())
    def domain(rel: str) -> str:
        if rel.startswith("基础档案/"):
            return "基础资料"
        if rel.startswith("单据数据/"):
            return "业务历史"
        if rel.startswith("枚举档案/"):
            return "枚举与尺码"
        if "库存" in rel:
            return "库存快照"
        return "其它来源"
    def domain_table(profile: dict) -> str:
        counted = Counter()
        for item in profile["byField"]:
            counted[(domain(item["file"]), item["issue"])] += item["count"]
        return "\n".join(f"| {area} | {issue} | {count:,} |" for (area, issue), count in sorted(counted.items()))
    (run / "CLEANING_BEFORE_PROFILE.md").write_text(f"# 清洗前全量扫描\n\n正式工作簿 {before['workbooks']}，Sheet {before['sheets']}，非空行 {before['nonemptyRows']:,}。\n\n| 候选类型 | 单元格数 |\n|---|---:|\n{table(before)}\n\n## 按数据域\n\n| 数据域 | 候选类型 | 单元格数 |\n|---|---|---:|\n{domain_table(before)}\n\n数字文本、全角字符、空字符串及备注换行是候选观察量，不等于可自动修复错误。\n")
    (run / "CLEANING_AFTER_PROFILE.md").write_text(f"# 清洗后全量扫描\n\n正式工作簿 {after['workbooks']}，Sheet {after['sheets']}，非空行 {after['nonemptyRows']:,}。\n\n| 候选类型 | 单元格数 |\n|---|---:|\n{table(after)}\n\n## 按数据域\n\n| 数据域 | 候选类型 | 单元格数 |\n|---|---|---:|\n{domain_table(after)}\n")
    before_counts = before["counts"]
    after_counts = after["counts"]
    rows = "\n".join(f"| {key} | {before_counts.get(key, 0):,} | {after_counts.get(key, 0):,} | {before_counts.get(key, 0) - after_counts.get(key, 0):,} |" for key in sorted(before_counts.keys() | after_counts.keys()))
    (run / "CLEANING_DIFF_REPORT.md").write_text(f"# 清洗前后差异\n\n| 扫描类型 | 前 | 后 | 减少 |\n|---|---:|---:|---:|\n{rows}\n\n实际修改单元格 {len(changes)}；人工复核队列 {len(queue)}；分类映射断言 {'PASS' if category['passed'] else 'STOP'}。原始文件值和字节未被本轮修改。\n")


def write_review_queue(run: Path, queue: list[dict]) -> None:
    with (run / "TPLUS_MANUAL_REVIEW_QUEUE.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(queue)
    (run / "TPLUS_MANUAL_REVIEW_QUEUE.md").write_text(f"# T+ 人工复核队列\n\n共 {len(queue)} 项，详见同目录 CSV。未填写 userDecision 的项目保留原值，不触发自动修复。\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--run", type=Path, required=True)
    parser.add_argument("--batch", required=True)
    parser.add_argument("--rules", type=Path, default=BASE / "TPLUS_CLEANING_MANIFEST.json")
    parser.add_argument("--structure", type=Path, default=BASE / "TPLUS_INPUT_STRUCTURE.json")
    parser.add_argument("--mapping", type=Path, default=BASE.parents[1] / "data-cleaning/authority/frozen-v1/decisions/categories/full_mapping_2941.json")
    parser.add_argument("--resume", action="store_true", help="verify and finish this batch's already written output")
    args = parser.parse_args()
    source = args.input.resolve()
    run = args.run.resolve()
    cleaned = run / "cleaned" / "T+数据"
    if source == cleaned or source in cleaned.parents or cleaned in source.parents:
        raise ValueError("raw and cleaned directories must be separate")
    manifest = json.loads(args.rules.read_text())
    structure = json.loads(args.structure.read_text())
    if manifest["inputStructureVersion"] != structure["inputStructureVersion"]:
        raise ValueError("manifest and input structure versions differ")
    if not args.batch.startswith("TPLUS-CLEAN-"):
        raise ValueError("cleaning batch must start with TPLUS-CLEAN-")
    run.mkdir(parents=True, exist_ok=True)
    existing_meta = run / "RUN_META.json"
    if existing_meta.exists():
        meta = json.loads(existing_meta.read_text())
        if meta["batch"] != args.batch or meta["rulesSha256"] != digest(args.rules) or meta["structureSha256"] != digest(args.structure):
            raise ValueError("existing cleaning batch has different rules or structure")
        if any(digest(source / rel) != sha for rel, sha in meta["rawHashes"].items()):
            raise ValueError("raw input changed since cleaning batch")
        if any(digest(cleaned / rel) != sha for rel, sha in meta["cleanedHashes"].items()):
            raise ValueError("cleaned output changed since cleaning batch")
        print(json.dumps({"status": "ALREADY_CLEAN", "batch": args.batch, "changedCells": meta["changedCells"]}))
        return
    formal, drift = validate_contract(source, structure)
    if drift:
        json_write(run / "INPUT_STRUCTURE_DRIFT.json", {"status": "INPUT_STRUCTURE_DRIFT", "differences": drift})
        raise ValueError(f"INPUT_STRUCTURE_DRIFT: {len(drift)} differences; see {run / 'INPUT_STRUCTURE_DRIFT.json'}")
    raw_hashes = {p.relative_to(source).as_posix(): digest(p) for p in sorted(source.rglob("*")) if p.is_file()}
    edits, changes, queue, category = collect(source, structure, manifest, args.mapping, args.batch)
    if not category["passed"]:
        write_review_queue(run, queue)
        json_write(run / "CATEGORY_ASSERTION_STOP.json", category)
    if cleaned.exists() and not args.resume:
        raise ValueError("cleaned output exists without completed RUN_META; keep for inspection")
    cleaned.mkdir(parents=True, exist_ok=True)
    for rel in sorted(raw_hashes):
        target = cleaned / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        if target.exists() and args.resume:
            continue
        if rel in edits:
            patch_workbook(source / rel, target, edits[rel], structure, rel)
        else:
            shutil.copy2(source / rel, target)
    cleaned_hashes = verify_output(source, cleaned, edits, raw_hashes)
    with (run / "cleaning_changes.jsonl").open("w") as handle:
        for change in changes:
            handle.write(json.dumps(change, ensure_ascii=False) + "\n")
    write_review_queue(run, queue)
    for name, source_dir in [("before-scan.json", source), ("after-scan.json", cleaned)]:
        subprocess.run([sys.executable, str(BASE / "scan_candidates.py"), "--input", str(source_dir), "--structure", str(args.structure), "--output", str(run / name)], check=True, stdout=subprocess.DEVNULL)
    render_profiles(run, changes, queue, category)
    meta = {
        "status": "CLEANED" if category["passed"] else "REVIEW_REQUIRED", "batch": args.batch, "ruleset": manifest["version"],
        "inputStructureVersion": structure["inputStructureVersion"],
        "input": str(source), "output": str(cleaned), "rulesSha256": digest(args.rules),
        "structureSha256": digest(args.structure), "mappingSha256": digest(args.mapping),
        "rawHashes": raw_hashes, "cleanedHashes": cleaned_hashes,
        "changedCells": len(changes), "manualReviewRows": len(queue),
        "categoryAssertion": category,
    }
    json_write(existing_meta, meta)
    print(json.dumps({"status": meta["status"], "batch": args.batch, "changedCells": len(changes), "manualReviewRows": len(queue), "output": str(cleaned)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
