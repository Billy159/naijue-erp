"""Read-only whole-folder value scan for conservative T+ cleaning candidates."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
NUMERIC = re.compile(r"^[+-]?(?:\d+\.\d+|\d+)$")
DATE = re.compile(r"^(?:19|20)\d\d[-/]\d{1,2}[-/]\d{1,2}$")
SENSITIVE = re.compile("密码|密钥|令牌|证件|手机|电话|邮箱|银行账号|^账号$|私钥|公钥", re.I)


def classify(value: str) -> list[str]:
    kinds = []
    if value != value.strip():
        kinds.append("outer_whitespace")
    if "\u00a0" in value:
        kinds.append("nbsp")
    if CTRL.search(value):
        kinds.append("control_character")
    if "\n" in value or "\r" in value or "\t" in value:
        kinds.append("embedded_linebreak_or_tab")
    if any("FULLWIDTH" in unicodedata.name(c, "") for c in value):
        kinds.append("fullwidth_character")
    if NUMERIC.fullmatch(value):
        kinds.append("numeric_text_candidate")
    if DATE.fullmatch(value):
        kinds.append("date_text_candidate")
    if value == "":
        kinds.append("empty_string")
    return kinds


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--structure", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    structure = json.loads(args.structure.read_text())
    if isinstance(structure, dict):
        definitions = {(item["file"], item["sheet"]): item for item in structure["sheets"]}
    else:
        definitions = {(Path(item["file"]).relative_to("T+数据").as_posix(), item["sheet"]): item for item in structure}
    counts = Counter()
    by_field = Counter()
    examples = defaultdict(list)
    sheets_done = 0
    rows_done = 0
    for path in sorted(args.input.rglob("*.xlsx")):
        if path.name.startswith(".~"):
            continue
        rel = path.relative_to(args.input).as_posix()
        book = load_workbook(path, read_only=True, data_only=False)
        for sheet in book:
            info = definitions[(rel, sheet.title)]
            header_row = info.get("header_row", info.get("headerRow"))
            data_start_row = info.get("dataStartRow", header_row + 1)
            headers = [col["field"] for col in info["fields"]] if "fields" in info else [str(x or "").lstrip("*").strip() for x in info["headers"]]
            for row_num, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_num < data_start_row or not any(v is not None and v != "" for v in values):
                    continue
                rows_done += 1
                for col_num, value in enumerate(values, 1):
                    if not isinstance(value, str):
                        continue
                    field = headers[col_num - 1] if col_num <= len(headers) else f"__column_{col_num}"
                    for kind in classify(value):
                        counts[kind] += 1
                        by_field[(kind, rel, sheet.title, field)] += 1
                        if len(examples[kind]) < 15:
                            examples[kind].append({
                                "file": rel, "sheet": sheet.title, "row": row_num,
                                "column": get_column_letter(col_num), "field": field,
                                "value": "[REDACTED]" if SENSITIVE.search(field) else value[:100],
                            })
            sheets_done += 1
            print(f"SCAN {sheets_done}/{len(definitions)} {rel} {sheet.title}", flush=True)
        book.close()
    report = {
        "workbooks": len({key[0] for key in definitions}), "sheets": sheets_done,
        "nonemptyRows": rows_done, "counts": counts,
        "byField": [
            {"issue": kind, "file": rel, "sheet": sheet, "field": field, "count": count}
            for (kind, rel, sheet, field), count in by_field.most_common()
        ],
        "examples": examples,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"workbooks": report["workbooks"], "sheets": sheets_done, "rows": rows_done, "counts": counts}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
