"""Read-only source extraction and deterministic preflight exchange helpers."""
from pathlib import Path
import os
from decimal import Decimal
from datetime import date, datetime
import csv, hashlib, json
from collections import Counter, defaultdict
from openpyxl import load_workbook

OUT = Path(os.environ['TPLUS_OUT_DIR']).resolve() if os.environ.get('TPLUS_OUT_DIR') else Path(__file__).resolve().parents[1]
ROOT = Path(os.environ['TPLUS_PROJECT_ROOT']).resolve() if os.environ.get('TPLUS_PROJECT_ROOT') else Path(__file__).resolve().parents[2]
SRC = Path(os.environ['TPLUS_SRC_DIR']).resolve() if os.environ.get('TPLUS_SRC_DIR') else ROOT / 'inputs/tplus-frozen-v1'
IDENTITY_SRC = Path(os.environ['TPLUS_IDENTITY_SRC_DIR']).resolve() if os.environ.get('TPLUS_IDENTITY_SRC_DIR') else SRC
AUDIT = ROOT / 'data-cleaning/authority/frozen-v1/audit'
DECISIONS = ROOT / 'data-cleaning/authority/frozen-v1/decisions'
_hashes = {}

def text(v):
    return '' if v is None else str(v)

def stable_id(kind, *parts):
    payload = json.dumps([text(x) for x in parts], ensure_ascii=False, separators=(',', ':'))
    return 'tp_' + kind + '_' + hashlib.sha256(payload.encode()).hexdigest()[:24]

def dec(v):
    return Decimal(str(v)) if v not in (None, '') else None

def decimal_text(v):
    d = v if isinstance(v, Decimal) else dec(v)
    if d is None:
        return None
    return format(d, 'f')

def serial(v):
    if isinstance(v, Decimal):
        return decimal_text(v)
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, set):
        return sorted(v)
    raise TypeError(type(v).__name__)

def dumps(v):
    return json.dumps(v, ensure_ascii=False, separators=(',', ':'), default=serial)

def sha(path):
    p = Path(path).resolve()
    if str(p) not in _hashes:
        h = hashlib.sha256()
        with p.open('rb') as f:
            for chunk in iter(lambda: f.read(1024*1024), b''):
                h.update(chunk)
        _hashes[str(p)] = h.hexdigest()
    return _hashes[str(p)]

def lineage(rel, sheet, row, **extra):
    p = SRC / rel
    return {'sourceFile': 'T+数据/' + rel, 'sourceSha256': sha(p), 'sheet': sheet,
            'row': row, **extra}

def raw_row_id(kind, rel, sheet, row):
    return stable_id(kind, 'T+数据/' + rel, sha(IDENTITY_SRC / rel), sheet, row)

def historical():
    return {'isHistorical': True, 'readOnly': True, 'auditStatus': None,
            'auditedBy': None, 'auditedAt': None, 'createdBy': None, 'createdAt': None,
            'sourceAuditStatusKnown': False, 'runtimeAuditReplayAllowed': False}

def source_rows(rel, sheet, header_row=1):
    """Yield (original row index, values, unique keyed fields), preserving duplicate headers."""
    book = load_workbook(SRC / rel, read_only=True, data_only=True)
    ws = book[sheet]
    counts = Counter(); headers = None
    try:
        for i, values in enumerate(ws.iter_rows(values_only=True), 1):
            if i < header_row:
                continue
            if i == header_row:
                headers = []
                for c, name in enumerate(values, 1):
                    name = text(name) or f'__column_{c}'
                    counts[name] += 1
                    headers.append(name if counts[name] == 1 else f'{name}__{counts[name]}')
                continue
            if any(v not in (None, '') for v in values):
                yield i, values, dict(zip(headers, values))
    finally:
        book.close()

def load_json(path):
    return json.loads(Path(path).read_text())

def jsonl(path):
    with Path(path).open() as f:
        for line in f:
            if line.strip():
                yield json.loads(line)

def save_json(rel, value):
    p = OUT / rel; p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=serial) + '\n')

def save_jsonl(rel, records):
    p = OUT / rel; p.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with p.open('w') as f:
        for r in records:
            f.write(dumps(r) + '\n'); count += 1
    return count

class JsonlWriter:
    def __init__(self, rel):
        self.path = OUT / rel; self.path.parent.mkdir(parents=True, exist_ok=True)
        self.file = self.path.open('w'); self.count = 0
    def write(self, obj):
        self.file.write(dumps(obj) + '\n'); self.count += 1
    def close(self):
        self.file.close()

class SpecResolver:
    def __init__(self):
        self.colors = defaultdict(list); self.sizes = defaultdict(list)
        for _, r, _ in source_rows('枚举档案/颜色.xlsx', '颜色'):
            self.colors[text(r[1])].append(text(r[0]))
        for _, r, _ in source_rows('枚举档案/尺码.xlsx', '尺码'):
            self.sizes[text(r[1])].append(text(r[0]))
    def resolve(self, product_code, color_name, size_name):
        p = text(product_code); cn = text(color_name); sn = text(size_name)
        cc = self.colors.get(cn, []); ss = self.sizes.get(sn, [])
        if len(cc) != 1 or len(ss) != 1:
            raise ValueError({'product': p, 'color': cn, 'size': sn, 'colorCandidates': cc, 'sizeCandidates': ss})
        return {'productId': stable_id('product', p), 'colorId': stable_id('color', cc[0]),
                'sizeId': stable_id('size', ss[0]), 'skuId': stable_id('sku', p, cc[0], ss[0]),
                'sourceProductCode': p, 'sourceColorCode': cc[0], 'sourceSizeCode': ss[0],
                'sourceColorName': cn, 'sourceSizeName': sn}

def control(entity, source_count, converted_count, quarantined_count=0, rejected_count=0,
            duplicate_count=0, missing_fk_count=0, **extra):
    return {'entity': entity, 'source_count': source_count, 'converted_count': converted_count,
            'rejected_count': rejected_count, 'quarantined_count': quarantined_count,
            'duplicate_count': duplicate_count, 'missing_fk_count': missing_fk_count, **extra}
