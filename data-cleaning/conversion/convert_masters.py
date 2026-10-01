"""Convert source masters into an isolated preflight exchange package. No runtime writes."""
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path
import ast, json, os, re
from openpyxl import load_workbook
from common import (ROOT, SRC, OUT, DECISIONS, stable_id, decimal_text, dec, text,
                    source_rows, lineage, sha, save_json, save_jsonl, load_json, control, dumps)

DEST='data/master/'
entities=[]; diagnostics=[]; missing_optional=[]; money_controls={}; outputs={}

def file_lineage(path, row=None, kind='confirmed-decision'):
    return {'sourceFile':str(path.relative_to(ROOT)), 'sourceSha256':sha(path),
            'sheet':None, 'row':row, 'kind':kind}

def source_base(rel,sheet,row,source):
    return {'origin':'tplus','lineage':lineage(rel,sheet,row),'source':source}

def nullable(v):
    return None if v in (None,'') else text(v)

def ref(kind,code,valid,entity,source_code,field,optional=True):
    if code in (None,''):
        if optional:missing_optional.append({'entity':entity,'sourceCode':source_code,'field':field,'reason':'source_null_preserved'})
        else:diagnostics.append({'entity':entity,'sourceCode':source_code,'field':field,'targetCode':None,'issue':'required_fk_missing'})
        return None
    code=text(code)
    if code not in valid:
        diagnostics.append({'entity':entity,'sourceCode':source_code,'field':field,'targetCode':code,'issue':'orphan_fk'})
        return None
    return stable_id(kind,code)

def emit(entity,records,source_count=None,converted_source_rows=None,**extra):
    source_count=len(records) if source_count is None else source_count
    ids=Counter(x['id'] for x in records);duplicates=sum(n-1 for n in ids.values() if n>1)
    missing=sum(d['entity']==entity for d in diagnostics)
    written=save_jsonl(DEST+entity+'.jsonl',records)
    assert written==len(records)
    outputs[entity]={'file':DEST+entity+'.jsonl','count':written,'importStatus':'可以导入','executed':False}
    entities.append(control(entity,source_count,written,duplicate_count=duplicates,
                            missing_fk_count=missing,
                            counting_grain=extra.pop('counting_grain',entity+' records'),
                            converted_source_rows=source_count if converted_source_rows is None else converted_source_rows,
                            **extra))
    assert not duplicates, (entity,ids)
    return records

def read(rel,sheet,header=1):return list(source_rows(rel,sheet,header))

CONFIRMED=ROOT/'data-cleaning/authority/frozen-v1/business-decisions/CONFIRMED_DECISIONS.md'
category_decision_path=DECISIONS/'categories/full_mapping_2941.json'
relation_path=DECISIONS/'relations-confirmed.json'
account_decision_path=DECISIONS/'accounts/account-decision-results.json'
category_map=load_json(category_decision_path);category_by={x['code']:x for x in category_map}
relations=load_json(relation_path);accounts_decision=load_json(account_decision_path)
assert len(category_map)==len(category_by)
if not os.environ.get('TPLUS_DYNAMIC_COUNTS'):
    assert len(category_map)==2941
assert all(x['target_id'] and not x['unresolved_candidates'] for x in category_map)
assert relations['counts']['warehouse_pending']==relations['counts']['customer_pending']==0
assert accounts_decision['summary']['scope_unresolved']==accounts_decision['summary']['primary_method_pending']==0

# Categories remain the existing topology; four already-existing parent IDs are assignable.
taxonomy_path=ROOT/'data-cleaning/authority/frozen-v1/frontend/product-master.js'
tree=ast.literal_eval(re.search(r'const tree=(.*?);',taxonomy_path.read_text()).group(1))
assignable_parents={'cat-8','cat-10','cat-11','cat-12'}
categories=[]
for i,(name,children) in enumerate(tree):
    parent_id='cat-'+str(i)
    categories.append({'id':parent_id,'name':name,'parentId':None,'order':i,
                       'isLeaf':not children,'allowsProductAssignment':not children or parent_id in assignable_parents,
                       'origin':'existing-confirmed-taxonomy','lineage':file_lineage(taxonomy_path,3,'runtime-taxonomy'),
                       'decisionLineage':file_lineage(CONFIRMED),
                       'source':{'name':name,'children':children,'existingId':parent_id}})
    for j,n in enumerate(children):
        categories.append({'id':parent_id+'-'+str(j),'name':n,'parentId':parent_id,'order':j,
                           'isLeaf':True,'allowsProductAssignment':True,'origin':'existing-confirmed-taxonomy',
                           'lineage':file_lineage(taxonomy_path,3,'runtime-taxonomy'),
                           'source':{'name':n,'existingId':parent_id+'-'+str(j)}})
emit('categories',categories,source_count=len(categories),derived_count=len(categories),source_kind='existing taxonomy nodes')
category_ids={x['id'] for x in categories if x['allowsProductAssignment']}
assert len(categories)==49 and sum(x['isLeaf'] for x in categories)==36 and len(category_ids)==40

dictionary={}
for entity,kind,file,sheet in [('colors','color','颜色','颜色'),('sizes','size','尺码','尺码'),('fits','fit','版型','版型')]:
    rel='枚举档案/'+file+'.xlsx';rows=read(rel,sheet)
    converted=[{'id':stable_id(kind,text(r['编码'])),'sourceCode':text(r['编码']),
                'code':text(r['编码']),'name':text(r['名称']),**source_base(rel,sheet,row,r)} for row,v,r in rows]
    dictionary[entity]=emit(entity,converted)
    assert len({r['code'] for r in converted})==len(converted)
size_names={r['name']:r for r in dictionary['sizes']};color_names={r['name']:r for r in dictionary['colors']}
fit_names={r['name']:r for r in dictionary['fits']}
assert len(size_names)==240 and len(color_names)==182 and len(fit_names)==11

department_rel='基础档案/部门.xlsx';dept_rows=read(department_rel,'部门')
dept_codes={text(r['部门编码']) for _,_,r in dept_rows}
hall_codes={x['code'] for x in accounts_decision['halls']}
assert len(hall_codes)==47
departments=[]
for row,v,r in dept_rows:
    code=text(r['部门编码'])
    departments.append({'id':stable_id('department',code),'sourceCode':code,'name':text(r['部门名称']),
                        'parentDepartmentId':ref('department',r['上级部门编码'],dept_codes,'departments',code,'parentDepartmentId'),
                        'hallId':stable_id('hall',code) if code in hall_codes else None,
                        'isHallSource':code in hall_codes,**source_base(department_rel,'部门',row,r)})
emit('departments',departments)

employee_rel='基础档案/员工.xlsx';staff_rows=read(employee_rel,'员工',2)
employee_codes={text(r['员工编码']) for _,_,r in staff_rows}
employees=[]
for row,v,r in staff_rows:
    code=text(r['员工编码']);dep=text(r['所属部门编码'])
    employees.append({'id':stable_id('employee',code),'sourceCode':code,'name':text(r['员工名称']),
                      'departmentId':ref('department',dep,dept_codes,'employees',code,'departmentId',False),
                      'hallId':stable_id('hall',dep) if dep in hall_codes else None,
                      'isSourceSalesperson':r['业务员']==1,
                      'loginIdentityCreated':False,'loginName':None,'loginEnabled':False,'role':None,
                      'authenticationMaterial':None,**source_base(employee_rel,'员工',row,r)})
emit('employees',employees)

party_rel='基础档案/往来单位.xlsx';party_rows=read(party_rel,'往来单位',2)
party_codes={text(r['往来单位编码']) for _,_,r in party_rows};parties=[]
for row,v,r in party_rows:
    code=text(r['往来单位编码']);nature=text(r['性质'])
    roles=[role for phrase,role in [('客户','customer'),('供应商','supplier')] if phrase in nature]
    assert roles
    parties.append({'id':stable_id('party',code),'sourceCode':code,'name':text(r['往来单位名称']),
                    'roles':roles,'tags':roles,'sourceNature':nature,
                    'managedDepartmentId':ref('department',r['分管部门编码'],dept_codes,'parties',code,'managedDepartmentId'),
                    'managedEmployeeId':ref('employee',r['分管人员编码'],employee_codes,'parties',code,'managedEmployeeId'),
                    'settlementCustomerId':ref('party',r['结算客户编码'],party_codes,'parties',code,'settlementCustomerId'),
                    **source_base(party_rel,'往来单位',row,r)})
emit('parties',parties)
for entity,role in [('customers','customer'),('suppliers','supplier')]:
    subset=[p for p in parties if role in p['roles']]
    emit(entity,[{'id':p['id'],'partyId':p['id'],'sourceCode':p['sourceCode'],'name':p['name'],
                  'role':role,'origin':'derived-party-role','lineage':p['lineage'],
                  'source':{'往来单位编码':p['sourceCode'],'性质':p['sourceNature']}} for p in subset],
         source_count=len(subset),derived_count=len(subset),input_party_rows=len(parties),
         counting_grain='one source party possessing the respective role')
supplier_codes={p['sourceCode'] for p in parties if 'supplier' in p['roles']}
customer_codes={p['sourceCode'] for p in parties if 'customer' in p['roles']}

warehouse_rel='基础档案/仓库.xlsx';warehouse_rows=read(warehouse_rel,'仓库')
warehouse_codes={text(r['仓库编码']) for _,_,r in warehouse_rows}
warehouse_decisions={r['code']:r for r in relations['warehouses']};warehouses=[]
for row,v,r in warehouse_rows:
    code=text(r['仓库编码']);decision=warehouse_decisions[code]
    shared=decision.get('related_departments',[]);single=decision['target_department']
    mode='rotating_shared' if shared else 'single_hall' if single else 'public'
    allowed=shared or ([single] if single else [])
    assert all(x in hall_codes for x in allowed)
    warehouses.append({'id':stable_id('warehouse',code),'sourceCode':code,'name':text(r['仓库名称']),
                       'hallId':stable_id('hall',single) if single else None,
                       'relatedHallIds':[stable_id('hall',x) for x in allowed],
                       'sourceHallCodes':allowed,'usageMode':mode,'rotationSchedule':None,
                       'historicalDepartmentRewriteAllowed':False,
                       'assignmentDecision':decision,'decisionLineage':file_lineage(relation_path),
                       **source_base(warehouse_rel,'仓库',row,r)})
emit('warehouses',warehouses)
customer_decisions={r['department']:r for r in relations['customers']};halls=[]
for d in departments:
    code=d['sourceCode']
    if code not in hall_codes:continue
    decision=customer_decisions[code];customer=decision['customer']['code']
    assigned=[w['id'] for w in warehouses if stable_id('hall',code) in w['relatedHallIds']]
    # Default warehouse is deliberately not inferred from the assigned warehouse list.
    halls.append({'id':stable_id('hall',code),'sourceCode':code,'name':d['name'],'departmentId':d['id'],
                  'defaultCustomerId':ref('party',customer,customer_codes,'halls',code,'defaultCustomerId',False),
                  'sourceDefaultCustomerCode':customer,'warehouseIds':assigned,'defaultWarehouseId':None,
                  'defaultWarehouseReason':'No explicit default warehouse decision in the confirmed input; no first-warehouse inference',
                  'origin':'derived-confirmed-hall','lineage':d['lineage'],'source':d['source'],
                  'assignmentDecision':decision,'decisionLineage':file_lineage(relation_path)})
emit('halls',halls,source_count=len(halls),derived_count=len(halls),input_department_rows=len(departments))

method_rel='基础档案/结算方式.xlsx';method_rows=read(method_rel,'结算方式')
method_codes={text(r['结算方式编码']) for _,_,r in method_rows};methods=[]
for row,v,r in method_rows:
    code=text(r['结算方式编码'])
    methods.append({'id':stable_id('payment_method',code),'sourceCode':code,'name':text(r['结算方式名称']),
                    'sourceDefaultAccountName':nullable(r['默认账号']),**source_base(method_rel,'结算方式',row,r)})
emit('payment_methods',methods)
derived_method={'id':stable_id('payment_method','mall-counter'),'sourceCode':None,'configurationKey':'mall-counter',
                'name':'商场收银台','origin':'existing-confirmed-derived-method',
                'lineage':file_lineage(ROOT/'data-cleaning/authority/frozen-v1/frontend/master-data.js',9,'runtime-configuration'),
                'source':{'existingConfigurationKey':'mall-counter','name':'商场收银台'},
                'decisionLineage':file_lineage(account_decision_path)}
emit('payment_methods_derived',[derived_method],source_count=1,derived_count=1,source_kind='existing runtime configuration')
method_codes.add('mall-counter')

# Financial account credentials (columns L:W) are not selected, inspected, or copied.
account_rel='基础档案/账号.xlsx';w=load_workbook(SRC/account_rel,read_only=True,data_only=True);ws=w['账号']
account_headers=next(ws.iter_rows(min_row=2,max_row=2,min_col=1,max_col=11,values_only=True))
disabled_values={i:r[0] for i,r in enumerate(ws.iter_rows(min_row=3,min_col=24,max_col=24,values_only=True),3)}
account_rows=[]
for row,v in enumerate(ws.iter_rows(min_row=3,min_col=1,max_col=11,values_only=True),3):
    if not any(x not in (None,'') for x in v):continue
    s=dict(zip(account_headers,v));s['停用']=disabled_values[row];account_rows.append((row,s))
w.close()
account_codes={text(r['账号编码']) for _,r in account_rows};account_by={r['code']:r for r in accounts_decision['records']}
fund_accounts=[]
for row,r in account_rows:
    code=text(r['账号编码']);decision=account_by[code];scope=decision['allowed_hall_source_codes'];method=decision['primary_method_id']
    assert method in method_codes and all(x in hall_codes for x in scope)
    source_evidence=dict(r);source_evidence['最新余额']=decimal_text(r['最新余额'])
    business_source={k:v for k,v in r.items() if k!='最新余额'}
    fund_accounts.append({'id':stable_id('fund_account',code),'sourceCode':code,'name':text(r['账号名称']),
                          'type':nullable(r['账号类型']),'bank':nullable(r['开户银行']),'accountNumber':nullable(r['账号']),
                          'sourceDepartmentId':ref('department',r['使用部门编码'],dept_codes,'fund_accounts',code,'sourceDepartmentId'),
                          'primaryPaymentMethodId':stable_id('payment_method',method),'sourcePrimaryMethodCode':method,
                          'hallIds':[stable_id('hall',x) for x in scope],'sourceHallCodes':scope,
                          'sourceEvidence':source_evidence,'balanceInitialized':False,
                          'credentialsImported':False,'excludedSourceColumnRange':'L:W',
                          'assignmentDecision':decision,'decisionLineage':file_lineage(account_decision_path),
                          **source_base(account_rel,'账号',row,business_source)})
emit('fund_accounts',fund_accounts)
money_controls['fund_accounts_source_latest_balance']={
    'source_sum':decimal_text(sum((dec(r['最新余额']) or Decimal(0) for _,r in account_rows),Decimal(0))),
    'converted_source_evidence_sum':decimal_text(sum((dec(r['sourceEvidence']['最新余额']) or Decimal(0) for r in fund_accounts),Decimal(0))),
    'negative_count':sum(dec(r['最新余额'])<0 for _,r in account_rows if dec(r['最新余额']) is not None),
    'operational_balance_initialized':False,'meaning':'source evidence only, no business opening balance'}

product_rel='基础档案/存货.xlsx';product_rows=read(product_rel,'存货',2)
product_codes={text(r['存货编码']) for _,_,r in product_rows}
assert product_codes==set(category_by)
price_fields=[('purchasePrice','采购价'),('price','零售价'),('wholesale5','五级批发价(工厂店供货价）'),
              ('wholesale6','六级批发价（贴牌价）'),('wholesale7','七级批发价（批发价）'),
              ('wholesale8','八级批发价（代理商供货价）'),('wholesale9','九级批发价（工厂店零售价）'),
              ('wholesale10','十级批发价（葛玉珠）'),('minimumPrice','最低售价')]
products=[];groups=defaultdict(list)
for row,v,r in product_rows:
    code=text(r['存货编码']);decision=category_by[code];target=decision['target_id']
    assert target in category_ids
    fit=fit_names.get(r['版型']);group=nullable(r['尺码所属取值组编码'])
    if group:groups[group].append((row,r))
    if r['版型'] not in (None,'') and fit is None:diagnostics.append({'entity':'products','sourceCode':code,'field':'fitId','issue':'unmatched_fit_name'})
    products.append({'id':stable_id('product',code),'sourceCode':code,'styleNumber':code,'name':text(r['存货名称']),
                     'categoryId':target,'categorySnapshotAtHistoricalTransaction':None,
                     'categoryDecision':decision,'categoryDecisionLineage':file_lineage(category_decision_path),
                     'unit':text(r['主计量']),'description':nullable(r['存货描述']),
                     'supplierId':ref('party',r['默认供应商编码'],supplier_codes,'products',code,'supplierId'),
                     'fitId':fit['id'] if fit else None,'fitCode':fit['code'] if fit else None,
                     'sourceSizeGroupId':stable_id('size_group',group) if group else None,
                     'sourceSizeGroupCode':group,'sourceColorGroupCode':nullable(r['颜色所属取值组编码']),
                     'completeAllowedColorSetKnown':False,'completeAllowedSkuSetKnown':False,
                     'defaultBarcode':text(r['默认条码']),'defaultBarcodeScope':'product_only',
                     'sourceBarcodeSchemeCode':text(r['条码方案识别码']),
                     'year':None,'season':None,'thickness':None,'images':[],
                     'sourceShapeName':nullable(r['领型或脚口']),
                     'defaultWarehouseId':ref('warehouse',r['默认仓库编码'],warehouse_codes,'products',code,'defaultWarehouseId'),
                     'newBarcodeGenerated':False,**{k:decimal_text(r[h]) for k,h in price_fields},
                     **source_base(product_rel,'存货',row,r)})
emit('products',products)
for field,source_field in price_fields:
    source_vals=[dec(r[source_field]) for _,_,r in product_rows];target_vals=[dec(r[field]) for r in products]
    assert source_vals==target_vals
    money_controls['products_'+field]={'source_sum':decimal_text(sum((v for v in source_vals if v is not None),Decimal(0))),
        'converted_sum':decimal_text(sum((v for v in target_vals if v is not None),Decimal(0))),
        'source_nonnull_count':sum(v is not None for v in source_vals),'converted_nonnull_count':sum(v is not None for v in target_vals),
        'null_count':sum(v is None for v in target_vals),'changed_value_count':0,'negative_count':sum(v is not None and v<0 for v in target_vals),
        'meaning':'reference price control, not transaction or cash total'}

size_groups=[]
for code,source_items in sorted(groups.items()):
    row,r=source_items[0]
    size_groups.append({'id':stable_id('size_group',code),'sourceCode':code,'name':text(r['尺码所属取值组']),
                        'origin':'derived-explicit-product-group-label','completeMembershipKnown':False,
                        'sizeIds':None,'lineage':lineage(product_rel,'存货',row),
                        'supportingLineage':[lineage(product_rel,'存货',i) for i,_ in source_items],
                        'source':{'尺码所属取值组编码':code,'尺码所属取值组':r['尺码所属取值组']}})
emit('size_groups',size_groups,source_count=len(groups),input_rows=len(product_rows),
     converted_source_rows=len(product_rows),derived_count=len(groups),counting_grain='distinct explicit source size-group labels')

barcode_rel='基础档案/存货条形码.xlsx';barcode_rows=read(barcode_rel,'存货条形码');barcodes=[]
for row,v,r in barcode_rows:
    product=text(r['存货编码']);barcode=text(r['条形码'])
    barcodes.append({'id':stable_id('product_barcode',product,barcode),'productId':ref('product',product,product_codes,'product_barcodes',barcode,'productId',False),
                     'sourceProductCode':product,'barcode':barcode,'unit':text(r['计量单位']),
                     'isDefault':r['是否默认条码']==1,'bindingScope':'product_only','skuId':None,
                     'skuBindingKnown':False,'generated':False,'isHistorical':True,'readOnly':True,
                     **source_base(barcode_rel,'存货条形码',row,r)})
emit('product_barcodes',barcodes)
unit_rel='基础档案/计量单位.xlsx';unit_rows=read(unit_rel,'计量单位')
emit('units',[{'id':stable_id('unit',text(r['计量单位名称'])),'name':text(r['计量单位名称']),
              **source_base(unit_rel,'计量单位',row,r)} for row,v,r in unit_rows])

# Verify complete source preservation and precise decisions in serialized output.
assert len(products)==len(product_rows)==len(category_map)
assert len({x['barcode'] for x in barcodes})==len(barcodes)
assert all(n<=1 for n in Counter(x['productId'] for x in barcodes if x['isDefault']).values())
if not os.environ.get('TPLUS_DYNAMIC_COUNTS'):
    assert len(products)==2941 and len(parties)==485 and len(employees)==99 and len(halls)==47
    assert len(warehouses)==51 and len(fund_accounts)==80 and len(barcodes)==2943
    assert sum(x['isDefault'] for x in barcodes)==2941
    assert sum(x['categoryId'] in assignable_parents for x in products)==143
    assert sum(len(x['hallIds'])==47 for x in fund_accounts)==24 and sum(len(x['hallIds'])==1 for x in fund_accounts)==56
    assert sum(x['assignmentDecision']['primary_method_status']=='user_confirmed' for x in fund_accounts)==21
shared=next(x for x in warehouses if x['sourceCode']=='96110')
assert shared['hallId'] is None and set(shared['sourceHallCodes'])=={'92101','91115'} and shared['rotationSchedule'] is None
assert next(x for x in warehouses if x['sourceCode']=='96102')['sourceHallCodes']==['91118']
assert next(x for x in halls if x['sourceCode']=='91118')['sourceDefaultCustomerCode']=='AC-05-93107'
assert not diagnostics,diagnostics
assert all(len(x['source'])==101 for x in products)
assert all('会员价' in x['source'] and '会员价__2' in x['source'] for x in products)
assert all('balance' not in x and '最新余额' not in x['source'] for x in fund_accounts)

readback={entity:[json.loads(line) for line in (OUT/meta['file']).read_text().splitlines()]
          for entity,meta in outputs.items()}
for entity,records in readback.items():
    assert len(records)==outputs[entity]['count']
    assert len({r['id'] for r in records})==len(records)
readback_products={r['sourceCode']:r for r in readback['products']}
for row,v,source in product_rows:
    converted=readback_products[text(source['存货编码'])]
    assert converted['source']==json.loads(dumps(source))
    assert converted['lineage']['row']==row
    for key,field in price_fields:assert converted[key]==decimal_text(source[field])

controls={'scope':'isolated preflight only; no runtime import','executed':False,'entities':entities,
          'moneyControls':money_controls,'fkDiagnostics':diagnostics,'missingOptionalSourceFields':missing_optional,
          'decisionControls':{'classificationAssigned':len(products),'classificationPending':0,'directParentAssignments':sum(x['categoryId'] in assignable_parents for x in products),
              'accountScopeAll47':sum(len(x['hallIds'])==len(halls) for x in fund_accounts),'accountScopeSingle':sum(len(x['hallIds'])==1 for x in fund_accounts),'newlyConfirmedOther991':sum(x['assignmentDecision']['primary_method_status']=='user_confirmed' for x in fund_accounts),'priorPrimaryMethodPreserved':sum(x['assignmentDecision']['primary_method_status']!='user_confirmed' for x in fund_accounts),
              'rotatingWarehouseSourceCode':'96110','rotationScheduleInvented':False},
          'sourcePreservation':{'productSourceColumns':101,'duplicateHeadersPreserved':True,
              'accountSelectedBusinessColumns':'A:K,X','accountCredentialColumnsNotSelected':'L:W',
              'accountLatestBalanceOnlyInSourceEvidence':True,'employeeLoginAccountsCreated':0,'skuRowsCreated':0},
          'verification':{'serializedFilesReadBack':len(readback),'sourceProductRowsCompared':len(product_rows),
              'sourceSerialization':'datetime values use the contract ISO representation; original fields and values otherwise preserved',
              'productPriceValuesCompared':len(product_rows)*len(price_fields),'allChecksPassed':True},
          'outputs':outputs,'inputDecisionLineages':[file_lineage(p) for p in [CONFIRMED,category_decision_path,relation_path,account_decision_path]],
          'conservationNotes':['Each direct source master has source_count = converted_count; rejected/quarantined = 0.',
              'Customer and supplier files are overlapping role projections of 485 parties, not additional parties.',
              f'{len(halls)} halls derive from {len(departments)} departments; {len(size_groups)} size groups derive from {len(products)} explicit product group labels.',
              'No color-by-size or warehouse-by-SKU Cartesian product is created.']}
save_json('controls/masters.json',controls)
notes='''# ③A-2 主档转换预检

全部输出只在本转换包 `data/master/`，`executed=false`；没有写当前 Demo、数据库、浏览器存储或原始 Excel。

采用最新 CONFIRMED_DECISIONS 和全2941分类映射，原901待确认均已收口。既有49分类节点（18一级、31子级）保持，36叶子之外4既有父级允许挂商品，本批143款父级直挂。所有商品保留原名、原条码、原描述、原价格、原来源；分类裁决不改物理属性，也不回填交易时分类。

## 输出与关系

'''
notes+='\n'.join(f"- `{e['entity']}.jsonl`：{e['converted_count']} 条；source_count={e['source_count']}，duplicate_count={e['duplicate_count']}，missing_fk_count={e['missing_fk_count']}。" for e in entities)
notes+='''

商品、party、employee、department、hall、warehouse、color、size、fit、payment_method、fund_account均使用 common.stable_id 与原源编码；分类沿用原 categoryId。customers/suppliers只是party角色投影，id和partyId沿用同一party，不制造第二套客户/供应商身份。

主档没有统一强制readOnly，便于后续合法维护；仅历史商品条码绑定证据标只读。styleNumber沿用原存货码只是单商品显示款号，不代表已识别同款多版型归并。

商品九价格是十进制字符串或null，逐值核对并保留控制合计，空值不补0。存货101列全部保留在source，两个同名“会员价”分别保留为会员价和会员价__2。year/season/thickness没有独立可信源字段，继续null，不因分类裁决填默认年季/版型。颜色组空不等于没有颜色或允许全部颜色；完整可选颜色/SKU集合标未知。SKU由主流程从真实观测组合生成，本脚本不生成矩阵。

主尺码240条只取“尺码”页，不用“Sheet1新尺码参考”覆盖910历史语义。版型11条按源编码保存，商品仅7个明确版型指向原01，未知不补A。商品条形码2943条只有productId，skuId=null，不能冒充完整可扫SKU绑定。

58部门保留职能部门历史外键，47营业厅由已确认关系派生。舟山96110保留一仓、两厅轮用关系，hallId=null、relatedHallIds两项、rotationSchedule=null；不编造月份和权限日历。庆春96102明确挂91118，91118默认客户为AC-05-93107。默认仓未在确认文件中明确，保持null，不选第一个仓。员工99条保留真实来源部门；不创建登录身份、口令和权限角色，原启停仅source留存。

账户80条使用最终已确认范围和唯一主要方式：24账户展开47厅，56单厅；21新裁决用991，其余59不变。11个历史方式和现有派生mall-counter分文件输出；派生项不混作第12条T+历史方式。每笔历史资金仍应保留原方式，本账户配置不覆盖交易方式。

账户只提取A:K、X业务字段；L:W支付凭证列不选择、不检查、不输出。最新余额只在sourceEvidence作为原值保留，43个负值保留，不建立业务余额、不参与库存或资金初始化。sourceEvidence控制合计-12903889.80为80个源值总和，不是当前可用资金。

## 核验

`controls/masters.json`含逐实体六计数、数量粒度、金额控制量、缺失可选字段明细及FK诊断。全量主档ID唯一，必需外键孤儿0、拒绝0、隔离0。缺失可选原字段保留null并解释，不计为错误外键。角色投影、分类和营业厅派生分别标注粒度，不能把这些数相加当新增业务实体总数。

所有直接来源记录有源文件SHA、Sheet、原行号；裁决及运行树派生项另有对应文件SHA与derivation/decision lineage。保留旧原始source状态不等于恢复旧启停流程。分类/账户/关系均按最新裁决执行预转换，未执行实际配置写入。
'''
(OUT/'notes').mkdir(exist_ok=True)
(OUT/'notes/MASTERS.md').write_text(notes)
print(json.dumps({'entities':[{k:e[k] for k in ['entity','source_count','converted_count','duplicate_count','missing_fk_count']} for e in entities],
                  'fkErrors':len(diagnostics),'classificationPending':0,'executed':False},ensure_ascii=False,indent=2))
