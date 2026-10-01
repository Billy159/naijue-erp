"""Receive a complete T+ folder and freeze its structure and accepted rules.

This entry has no database dependency and performs no business-data conversion.
"""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def ignored(path):
    return any(p == '.DS_Store' or p.startswith(('~$', '.~', '._')) for p in path.parts)


def inventory(source):
    if not source.is_dir():
        raise ValueError('T+输入文件夹不存在')
    files, skipped = [], []
    for path in sorted(source.rglob('*')):
        rel = path.relative_to(source)
        if ignored(rel):
            if path.is_file():
                skipped.append(rel.as_posix())
            continue
        if path.is_symlink():
            raise ValueError('输入含符号链接，请提供实际原件：' + rel.as_posix())
        if path.is_file():
            files.append({'path': rel.as_posix(), 'bytes': path.stat().st_size, 'sha256': digest(path)})
    if not files:
        raise ValueError('T+输入文件夹为空')
    return sorted(files, key=lambda item: item['path']), sorted(skipped)


def headers(path, definitions):
    """Only read approved header rows; never read account/password data rows."""
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        result = []
        for definition in definitions:
            name = definition['sheet']
            if name not in book.sheetnames:
                continue
            sheet = book[name]
            rows = []
            for row_number in definition['headerRows']:
                values = list(next(sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True)))
                while values and values[-1] is None:
                    values.pop()
                rows.append({'row': row_number, 'values': values})
            result.append({'sheet': name, 'headers': rows})
        return {'sheetNames': book.sheetnames, 'sheets': result}
    finally:
        book.close()


def check_structure(source, files, contract):
    actual = {f['path'] for f in files}
    expected = {f['path']: f for f in contract['files']}
    problems, observations = [], []
    for rel, item in expected.items():
        if rel not in actual:
            if item['required']:
                problems.append({'kind': 'MISSING_FILE', 'path': rel})
            else:
                observations.append({'kind': 'MISSING_REFERENCE', 'path': rel})
    for rel in sorted(actual):
        item = expected.get(rel)
        if item is None:
            problems.append({'kind': 'NEW_FILE_REQUIRES_MAPPING', 'path': rel})
            continue
        if item['kind'] != 'workbook':
            continue
        try:
            observed = headers(source / rel, item['sheets'])
        except Exception as error:
            problems.append({'kind': 'UNREADABLE_WORKBOOK', 'path': rel, 'errorType': type(error).__name__})
            continue
        if observed != item['structure']:
            problems.append({'kind': 'SHEET_OR_HEADER_CHANGED', 'path': rel, 'expected': item['structure'], 'actual': observed})
    return problems, observations


def rules_inventory(project, registry):
    files = []
    for rel in registry['files']:
        path = project / rel
        if not path.is_file():
            raise ValueError('已裁定规则依赖缺失：' + rel)
        files.append({'path': rel, 'bytes': path.stat().st_size, 'sha256': digest(path)})
    return files


def copy_verified(source, destination, expected_sha):
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        if not destination.is_file() or digest(destination) != expected_sha:
            raise ValueError('已有批次副本内容不一致：' + str(destination))
    else:
        shutil.copy2(source, destination)
    if digest(source) != expected_sha or digest(destination) != expected_sha:
        raise ValueError('接收过程中原件或副本发生变化：' + str(source))


def verify_receipt(run, files, rules):
    for name, expected in [('source', files), ('rules', rules)]:
        actual, skipped = inventory(run / name)
        if actual != sorted(expected, key=lambda item: item['path']) or skipped:
            raise ValueError('接收批次副本文件集合或内容发生变化：' + name)
    for item in files:
        path = run / 'source' / item['path']
        if not path.is_file() or digest(path) != item['sha256']:
            raise ValueError('接收批次原件副本缺失或变化：' + item['path'])
    for item in rules:
        path = run / 'rules' / item['path']
        if not path.is_file() or digest(path) != item['sha256']:
            raise ValueError('接收批次裁定副本缺失或变化：' + item['path'])


def receive(source, run, contract, registry, project=ROOT, cutoff=None):
    source = source.resolve()
    run = run.resolve()
    if source == run or source in run.parents or run in source.parents:
        raise ValueError('接收输出与输入目录不能重叠')
    if cutoff is not None:
        value = datetime.fromisoformat(cutoff)
        if value.tzinfo is None:
            raise ValueError('库存截止时点必须包含时区，例如 +08:00')
    files, skipped = inventory(source)
    rules = rules_inventory(project, registry)
    source_sha = hashlib.sha256(canonical(files)).hexdigest()
    rule_sha = hashlib.sha256(canonical(rules)).hexdigest()
    contract_sha = hashlib.sha256(canonical(contract)).hexdigest()
    identity = {'sourceFingerprint': source_sha, 'rulesFingerprint': rule_sha, 'contractFingerprint': contract_sha, 'cutoff': cutoff}
    receipt = run / 'RECEIPT.json'
    if receipt.exists():
        saved = json.loads(receipt.read_text())
        if saved['identity'] != identity:
            raise ValueError('同一接收目录的原件、裁定、结构合同或截止时点已改变，请使用新批次目录')
        verify_receipt(run, files, rules)
        if inventory(source)[0] != files or rules_inventory(project, registry) != rules:
            raise ValueError('校验过程中原件或裁定发生变化')
        return {**saved, 'repeatStatus': 'ALREADY_RECEIVED'}
    problems, observations = check_structure(source, files, contract)
    run.mkdir(parents=True, exist_ok=True)
    for item in files:
        copy_verified(source / item['path'], run / 'source' / item['path'], item['sha256'])
    for item in rules:
        copy_verified(project / item['path'], run / 'rules' / item['path'], item['sha256'])
    if inventory(source)[0] != files or rules_inventory(project, registry) != rules:
        raise ValueError('接收结束时原件或裁定发生变化，未确认此批次')
    verify_receipt(run, files, rules)
    result = {
        'version': 'TPLUS-FOLDER-RECEIPT-v1', 'mode': 'FULL_REBUILD_FROM_FOLDER',
        'status': 'RECEIVED' if not problems else 'REVIEW_REQUIRED',
        'receivedAt': datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),
        'sourceDirectory': str(source), 'batchDirectory': str(run), 'identity': identity,
        'files': files, 'ignoredTemporaryFiles': skipped, 'ruleFiles': rules,
        'summary': {'files': len(files), 'workbooks': sum(f['path'].endswith('.xlsx') for f in files),
                    'referenceFiles': sum(not f['path'].endswith('.xlsx') for f in files),
                    'bytes': sum(f['bytes'] for f in files), 'ruleFiles': len(rules),
                    'structureIssues': len(problems), 'observations': len(observations)},
        'structureIssues': problems, 'observations': observations,
        'originalFilesUnchanged': True, 'ruleFilesUnchanged': True, 'databaseWrites': 0,
        'businessCleaningExecuted': False, 'databaseBuildExecuted': False,
        'openingPolicy': 'CONFIRMED_FULL_WAREHOUSE_SNAPSHOT',
        'cutoffStatus': 'PROVIDED_FOR_REVIEW' if cutoff else 'NOT_CONFIRMED',
        'accountPolicy': 'PRESERVE_HISTORY_IDENTITIES_LOGIN_REORGANIZATION_UNDECIDED',
        'nextStage': '清洗和来源身份匹配，再在独立新库导入与核对；正式使用前确认完整全仓库存及统一截止时点',
    }
    write_json(run / 'INPUT_FILES.json', files)
    write_json(run / 'ACCEPTED_RULES.json', {'registry': registry, 'files': rules})
    write_json(run / 'STRUCTURE_REVIEW.json', {'issues': problems, 'observations': observations})
    s = result['summary']
    lines = ['# T+数据文件夹接收结果', '', f"状态：{result['status']}。接收时间：{result['receivedAt']}。", '',
             f"已保存 {s['workbooks']} 份工作簿、{s['referenceFiles']} 份参考文件和 {s['ruleFiles']} 份裁定 / 规则依赖。",
             f"结构问题 {s['structureIssues']} 项。忽略 WPS / Excel 临时文件 {len(skipped)} 个。原件与裁定副本逐份校验一致。", '',
             '后续方式：整套新T+数据配合既定裁定，在独立空数据库重建。',
             '本次完成原件接收、结构检查和规则封存，尚未清洗业务行或建立数据库。',
             '库存使用本批完整全仓快照；共同截止时点待确认。历史单据不重复叠加期初。',
             '人员历史身份保留，登录账号是否重整尚未裁定。本次未读取账号密码数据行。', '',
             '## 检查事项', '']
    lines += [f"- {p['kind']}：{p['path']}" for p in problems] or ['当前文件结构与接收模板一致。']
    if observations:
        lines += ['', '## 参考资料差异', ''] + [f"- {p['kind']}：{p['path']}" for p in observations]
    lines += ['', '完整清单和核对记录见 RECEIPT.json、INPUT_FILES.json、ACCEPTED_RULES.json 与 STRUCTURE_REVIEW.json。',
              'source/保存本批原件；rules/保存本批采用的规则版本。']
    (run / 'README.md').write_text('\n'.join(lines) + '\n')
    write_json(receipt, result)
    return result


def main():
    parser = argparse.ArgumentParser(description='接收整套T+数据，检查结构、封存原件与裁定；不写数据库')
    parser.add_argument('--input', type=Path, default=ROOT / 'inputs/tplus-current')
    parser.add_argument('--run', type=Path)
    parser.add_argument('--cutoff', help='用户提供的共同库存截止时点，需包含时区')
    args = parser.parse_args()
    contract = json.loads((HERE / 'input-contract.json').read_text())
    registry = json.loads((HERE / 'decision-registry.json').read_text())
    if args.run:
        run = args.run
    else:
        files, _ = inventory(args.input.resolve())
        rules = rules_inventory(ROOT, registry)
        identifier = hashlib.sha256(canonical({'files': files, 'rules': rules, 'contract': contract, 'cutoff': args.cutoff})).hexdigest()[:16]
        run = ROOT / 'data-cleaning/runs/received' / identifier
    result = receive(args.input, run, contract, registry, cutoff=args.cutoff)
    print(json.dumps({'status': result['status'], 'repeatStatus': result.get('repeatStatus'),
                      'batchDirectory': result['batchDirectory'], **result['summary'], 'databaseWrites': 0}, ensure_ascii=False))
    raise SystemExit(0 if result['status'] == 'RECEIVED' else 2)


if __name__ == '__main__':
    main()
