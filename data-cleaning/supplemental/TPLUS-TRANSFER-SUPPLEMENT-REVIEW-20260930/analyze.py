"""Read-only review of supplemental transfers and 131 quarantined source documents."""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
from datetime import datetime, date
import hashlib, json, re, subprocess
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
WORK = Path(__file__).resolve().parent
SOURCE = ROOT / 'inputs/tplus-current'
CUTOFF = '2026-09-10'

def text(v):
    return '' if v is None else str(v)

def decimal(v):
    return Decimal(text(v) or '0')

def json_default(v):
    if isinstance(v, Decimal): return str(v)
    if isinstance(v, (datetime,date)): return v.isoformat()
    if isinstance(v, (set, tuple)): return list(v)
    raise TypeError(type(v).__name__)

def write(name, value):
    (WORK / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=json_default)+'\n')

def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024*1024), b''): h.update(chunk)
    return h.hexdigest()

def item_key(product, color, size, quantity):
    return text(product), text(color), text(size), decimal(quantity)

def read_old(rel, side):
    path = SOURCE / rel
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    docs = {}
    headers = None
    count = 0
    for row, vals in enumerate(sheet.iter_rows(values_only=True), 1):
        if row == 2: headers = list(vals)
        if row < 3 or not any(v not in (None, '') for v in vals): continue
        number = text(vals[1])
        assert number, (rel, row, 'empty number')
        if number not in docs:
            docs[number] = {'number':number, 'side':side, 'date':text(vals[0])[:10],
                'typeCode':text(vals[2]), 'type':text(vals[3]), 'category':text(vals[5]),
                'note':text(vals[12]), 'warehouseCode':text(vals[13]), 'warehouse':text(vals[14]),
                'row':row, 'lastRow':row, 'lineCount':0, 'quantity':Decimal(0),
                'amount':Decimal(0), 'units':Counter(), 'positiveRows':0, 'negativeRows':0,
                'zeroRows':0, 'details':Counter(), 'skuQty':defaultdict(Decimal),
                'header':dict(zip(headers, vals[:17])), 'nonemptyNotes':set()}
        d = docs[number]
        q = decimal(vals[23])
        d['lineCount'] += 1; d['lastRow'] = row; d['quantity'] += q
        d['amount'] += decimal(vals[24]); d['units'][text(vals[22])] += 1
        d['positiveRows' if q > 0 else 'negativeRows' if q < 0 else 'zeroRows'] += 1
        d['details'][item_key(vals[17], vals[20], vals[21], q)] += 1
        d['skuQty'][(text(vals[17]), text(vals[20]), text(vals[21]))] += q
        if vals[12]: d['nonemptyNotes'].add(text(vals[12]))
        count += 1
    book.close()
    print(json.dumps({'read':rel, 'documents':len(docs), 'lines':count}, ensure_ascii=False), flush=True)
    return docs, {'path':str(path), 'sheet':sheet.title, 'sha256':sha(path), 'documents':len(docs), 'lines':count}

def read_new():
    path = SOURCE / '调拨单导出.xlsx'
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book['调拨单']
    docs = {}; current = None; detail_count = 0; header_count = 0; repeated = []
    for row, vals in enumerate(sheet.iter_rows(values_only=True), 1):
        if row == 1: continue
        if not any(v not in (None, '') for v in vals): continue
        if vals[1]:
            header_count += 1
            number = text(vals[1])
            assert number not in docs, (number,row,'duplicate transfer header')
            current = {'number':number,'date':text(vals[2])[:10], 'from':text(vals[3]), 'to':text(vals[4]),
                'type':text(vals[0]), 'note':text(vals[5]), 'creator':text(vals[6]), 'auditor':text(vals[7]),
                'auditDate':text(vals[8]), 'status':text(vals[9]), 'row':row, 'lastRow':row,
                'lineCount':0, 'quantity':Decimal(0), 'details':Counter(), 'skuQty':defaultdict(Decimal)}
            docs[number] = current
        assert current is not None, (row,'orphan line')
        if vals[11]:
            assert vals[14] is not None, (row,'missing quantity')
            q = decimal(vals[14]); current['quantity'] += q; current['lineCount'] += 1
            current['details'][item_key(vals[11],vals[12],vals[13],q)] += 1
            current['skuQty'][(text(vals[11]),text(vals[12]),text(vals[13]))] += q
            current['lastRow'] = row; detail_count += 1
    book.close()
    print(json.dumps({'read':path.name,'documents':len(docs),'lines':detail_count}, ensure_ascii=False),flush=True)
    return docs, {'path':str(path),'sheet':sheet.title,'sha256':sha(path),'documents':len(docs),
        'lines':detail_count, 'minDate':min(d['date'] for d in docs.values()),
        'maxDate':max(d['date'] for d in docs.values())}

def serialize_doc(d):
    return {k:v for k,v in d.items() if k not in ('details','skuQty','units','nonemptyNotes')} | {
        'units':dict(d.get('units',{})), 'notes':sorted(d.get('nonemptyNotes',set()))}

def sku_fingerprint(d):
    pairs = sorted((p,c,s,str(q.normalize())) for (p,c,s),q in d['skuQty'].items() if q != 0)
    return hashlib.sha256(json.dumps(pairs,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()

def compare_skus(a,b):
    keys = set(a['skuQty']) | set(b['skuQty'])
    diff = [{'product':p,'color':c,'size':s,'outQty':a['skuQty'].get((p,c,s),Decimal(0)),
             'inQty':b['skuQty'].get((p,c,s),Decimal(0)),
             'net':b['skuQty'].get((p,c,s),Decimal(0))-a['skuQty'].get((p,c,s),Decimal(0))}
        for p,c,s in sorted(keys) if a['skuQty'].get((p,c,s),Decimal(0)) != b['skuQty'].get((p,c,s),Decimal(0))]
    return {'equalNonzeroSkuQuantities':not diff,'differentSkuKeys':len(diff),'differences':diff}

old_out, out_info = read_old('单据数据/其他出库单.xlsx','out')
old_in, in_info = read_old('单据数据/其他入库单.xlsx','in')
new, new_info = read_new()
common = set(old_out) & set(old_in)
qout = {n:old_out[n] for n in sorted(set(old_out)-set(old_in))}
qin = {n:old_in[n] for n in sorted(set(old_in)-set(old_out))}
assert (len(common),len(qout),len(qin)) == (13946,52,79)
assert (sum(d['lineCount'] for d in qout.values()),sum(d['lineCount'] for d in qin.values())) == (44402,44398)
assert (sum(d['quantity'] for d in qout.values()),sum(d['quantity'] for d in qin.values())) == (50050,51645)
old_dates = [d['date'] for d in old_out.values()] + [d['date'] for d in old_in.values()]
assert max(old_dates) == CUTOFF
legacy_new = {n:d for n,d in new.items() if d['date'] <= CUTOFF}
future_new = {n:d for n,d in new.items() if d['date'] > CUTOFF}

index = defaultdict(list)
for d in legacy_new.values():
    for side in ('from','to'):
        index[(d['date'],side,d[side],sku_fingerprint(d))].append(d)
q_matches=[]
all_q = list(qout.values())+list(qin.values())
for d in all_q:
    side = 'from' if d['side']=='out' else 'to'
    exact_ids = [n for n in legacy_new if n == d['number']]
    linked_notes = [n for n,t in legacy_new.items() if d['number'] in t['note']]
    candidates = index[(d['date'],side,d['warehouse'],sku_fingerprint(d))]
    verified_candidates = [c['number'] for c in candidates if compare_skus(d,c)['equalNonzeroSkuQuantities']]
    q_matches.append(serialize_doc(d)|{'sameNumberTransfers':exact_ids,'mentionedInTransferNotes':linked_notes,
        'sameDateWarehouseAndCompleteNonzeroSkuQuantityTransfers':verified_candidates})

# Compare reliable original transfers with the supplemented direct transfer export.
transfer_index = defaultdict(list)
for n in common:
    a,b = old_out[n],old_in[n]
    transfer_index[(a['date'],a['warehouse'],b['warehouse'],sku_fingerprint(a))].append(n)
existing_matches=[]; unmatched_legacy=[]; ambiguity=[]
for n,d in legacy_new.items():
    candidate=transfer_index[(d['date'],d['from'],d['to'],sku_fingerprint(d))]
    good=[old for old in candidate if compare_skus(old_out[old],d)['equalNonzeroSkuQuantities']]
    if good:
        existing_matches.append({'newNumber':n,'row':d['row'],'oldNumbers':good,'status':d['status']})
        if len(good)>1: ambiguity.append({'newNumber':n,'oldNumbers':good})
    else: unmatched_legacy.append(serialize_doc(d))

candidate_pairs=[]
for n,a in qout.items():
    companion = 'IC-' + n.removeprefix('ID-')
    if companion not in qin: continue
    b=qin[companion]
    same_wh=a['warehouseCode']==b['warehouseCode']; same_date=a['date']==b['date']
    comparison=compare_skus(a,b)
    candidate_pairs.append({'out':n,'in':companion,'sameWarehouse':same_wh,'sameDate':same_date,
        'outWarehouse':a['warehouse'],'inWarehouse':b['warehouse'],'outDate':a['date'],'inDate':b['date'],
        'outQuantity':a['quantity'],'inQuantity':b['quantity'],'outNote':a['note'],'inNote':b['note'],
        'equalNonzeroSkuQuantities':comparison['equalNonzeroSkuQuantities'],
        'differentSkuKeys':comparison['differentSkuKeys']})

special_a=qout['ID-2024-08-0006']; special_b=qin['IC-2024-08-0006']
special= {'out':serialize_doc(special_a),'in':serialize_doc(special_b),**compare_skus(special_a,special_b)}
stocktaking_notes=[serialize_doc(d) for d in legacy_new.values() if re.search('盘点|盘盈|盘亏|IC-20|ID-20', d['note'])]
summary={'cutoff':CUTOFF,'inputs':[out_info,in_info,new_info],
    'newExport':{'documents':len(new),'lines':new_info['lines'],'legacyDocuments':len(legacy_new),
        'ignoredNewDocuments':len(future_new),'ignoredNewLines':sum(d['lineCount'] for d in future_new.values()),
        'allStatuses':dict(Counter(d['status'] for d in new.values())),
        'legacyStatuses':dict(Counter(d['status'] for d in legacy_new.values())),
        'legacyTypes':dict(Counter(d['type'] for d in legacy_new.values()))},
    'quarantine':{'documents':len(all_q),'bySide':{'out':len(qout),'in':len(qin)},
        'lines':sum(d['lineCount'] for d in all_q),'byType':dict(Counter(d['type'] for d in all_q)),
        'outQuantity':sum(d['quantity'] for d in qout.values()),'inQuantity':sum(d['quantity'] for d in qin.values()),
        'arithmeticInMinusOut':sum(d['quantity'] for d in qin.values())-sum(d['quantity'] for d in qout.values()),
        'dateMin':min(d['date'] for d in all_q),'dateMax':max(d['date'] for d in all_q),
        'notes':dict(Counter(d['note'] for d in all_q)),
        'exactNumberMatches':sum(bool(d['sameNumberTransfers']) for d in q_matches),
        'noteMentions':sum(bool(d['mentionedInTransferNotes']) for d in q_matches),
        'completeSameDateWarehouseSkuQuantityMatches':sum(bool(d['sameDateWarehouseAndCompleteNonzeroSkuQuantityTransfers']) for d in q_matches)},
    'originalTransfers':{'documents':len(common),'newDocsWithExactDateWarehousesSkuQuantities':len(existing_matches),
        'matchedOldDocuments':len({old for e in existing_matches for old in e['oldNumbers']}),
        'legacyNewDocsWithoutExactMatch':len(unmatched_legacy),'ambiguity':len(ambiguity)},
    'sameMonthlySerialIcIdCandidates':{'pairs':len(candidate_pairs),
        'sameDateAndWarehouse':sum(p['sameWarehouse'] and p['sameDate'] for p in candidate_pairs),
        'sameTotalQuantity':sum(p['outQuantity']==p['inQuantity'] for p in candidate_pairs),
        'sameCompleteNonzeroSkuQuantity':sum(p['equalNonzeroSkuQuantities'] for p in candidate_pairs)},
    'notesMentioningStocktaking':len(stocktaking_notes),'readOnly':True,'databaseWritten':False}
write('summary.json',summary)
write('quarantined-131-documents.json',q_matches)
write('original-transfer-matches.json',existing_matches)
write('legacy-new-without-exact-match.json',unmatched_legacy)
write('legacy-unapproved-transfers.json',[serialize_doc(d) for d in legacy_new.values() if d['status']!='生效'])
number_identity_comparisons=[]
for n,d in legacy_new.items():
    old_number = 'AL' + n
    if old_number not in common:
        number_identity_comparisons.append({'newNumber':n,'oldCandidate':old_number,'issue':'no original paired transfer with this serial'})
        continue
    a,b=old_out[old_number],old_in[old_number]
    comp=compare_skus(a,d)
    same_headers=(d['date'],d['from'],d['to']) == (a['date'],a['warehouse'],b['warehouse'])
    if not same_headers or not comp['equalNonzeroSkuQuantities'] or d['status']!='生效':
        number_identity_comparisons.append({'new':serialize_doc(d),'oldOut':serialize_doc(a),
            'oldIn':serialize_doc(b),'sameDateAndWarehouses':same_headers,'comparison':comp})
write('same-serial-transfer-differences-and-status.json',number_identity_comparisons)
write('transfer-matching-ambiguity.json',ambiguity)
write('same-monthly-number-ic-id-candidates.json',candidate_pairs)
write('ID-IC-2024-08-0006-comparison.json',special)
write('transfer-notes-mentioning-stocktaking.json',stocktaking_notes)
print(json.dumps(summary,ensure_ascii=False,default=json_default),flush=True)
print('SPECIAL',json.dumps({k:v for k,v in special.items() if k!='differences'},ensure_ascii=False,default=json_default),flush=True)
